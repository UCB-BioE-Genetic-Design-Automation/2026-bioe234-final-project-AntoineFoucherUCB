from __future__ import annotations

import asyncio
import ast
import base64
import json
import os
import shutil
import subprocess
import tempfile
import threading
from pathlib import Path
from typing import Any, Coroutine, Optional, TypeVar

import streamlit as st
from streamlit.errors import StreamlitAPIException
from docx import Document
from dotenv import load_dotenv
from streamlit.components.v1 import html as st_html

from mcp_gemini_engine import EngineState, MCPGeminiEngine

T = TypeVar("T")

PROJECT_DIR = Path(__file__).parent
_LOOP_LOCK = threading.Lock()
_BG_LOOP: Optional[asyncio.AbstractEventLoop] = None


def _get_background_loop() -> asyncio.AbstractEventLoop:
    global _BG_LOOP
    with _LOOP_LOCK:
        if _BG_LOOP is not None and _BG_LOOP.is_running():
            return _BG_LOOP
        loop = asyncio.new_event_loop()

        def _run() -> None:
            asyncio.set_event_loop(loop)
            loop.run_forever()

        t = threading.Thread(target=_run, daemon=True, name="streamlit-mcp-loop")
        t.start()
        _BG_LOOP = loop
        return _BG_LOOP


def run_coroutine_sync(coro: Coroutine[Any, Any, T], *, timeout: float = 90.0) -> T:
    loop = _get_background_loop()
    fut = asyncio.run_coroutine_threadsafe(coro, loop)
    return fut.result(timeout=timeout)


def _get_engine() -> MCPGeminiEngine:
    if "engine" not in st.session_state:
        _load_env_files(PROJECT_DIR)
        st.session_state.engine = MCPGeminiEngine(project_dir=PROJECT_DIR)
    return st.session_state.engine


def _init_state() -> None:
    st.session_state.setdefault("messages", [])
    st.session_state.setdefault("last_doc_path", None)
    st.session_state.setdefault("last_render_text", None)
    st.session_state.setdefault("live_bua_state", None)
    st.session_state.setdefault("mode", None)
    st.session_state.setdefault("uploaded_path", None)
    st.session_state.setdefault("upload_processed", False)
    st.session_state.setdefault("questionnaire_complete", False)
    st.session_state.setdefault("questionnaire_started", False)


async def _send(engine: MCPGeminiEngine, user_text: str):
    return await engine.send_message(user_text)


def _load_env_files(project_dir: Path) -> None:
    custom_file = os.environ.get("BIOE_ENV_FILE", "").strip()
    if custom_file:
        p = Path(custom_file)
        if p.exists():
            load_dotenv(p, override=False)
        return

    candidates = [
        project_dir / "ABSA_KEY",
        project_dir.parent / "ABSA_KEY",
        project_dir / ".env.gemini",
        project_dir.parent / ".env.gemini",
        project_dir / ".env",
        project_dir.parent / ".env",
        project_dir / ".env.local",
        project_dir.parent / ".env.local",
    ]
    for env_file in candidates:
        if env_file.exists():
            load_dotenv(env_file, override=False)


def _render_debug_panel(state: Optional[EngineState]) -> None:
    if state is None:
        return
    with st.expander("Debug details", expanded=False):
        st.write("Last generated doc path:", state.last_generated_doc_path)
        if state.last_tool_events:
            st.write("Last tool calls:")
            for ev in state.last_tool_events[-8:]:
                st.json(
                    {
                        "tool_name": ev.tool_name,
                        "args": ev.args,
                        "response_preview": str(ev.response)[:500],
                    }
                )


@st.cache_data(show_spinner=False)
def _load_tool_library() -> list[dict[str, str]]:
    tools_dir = PROJECT_DIR / "modules" / "biosafety" / "tools"
    rows: list[dict[str, str]] = []
    for p in sorted(tools_dir.glob("*.json")):
        try:
            data = json.loads(p.read_text(encoding="utf-8"))
        except Exception:
            continue
        name = str(data.get("name", p.stem))
        desc = str(data.get("description", "")).strip()
        rows.append({"name": name, "description": desc})
    return rows


def _render_sidebar_tool_library() -> None:
    with st.sidebar:
        with st.expander("tool library ⚙️", expanded=False):
            for item in _load_tool_library():
                st.markdown(f"**{item['name']}**")
                if item["description"]:
                    st.caption(item["description"])
                st.markdown("---")


def _parse_tool_result_payload(raw: Any) -> dict[str, Any]:
    if isinstance(raw, dict):
        return raw
    if not isinstance(raw, str):
        return {}
    txt = raw.strip()
    if not txt:
        return {}
    try:
        parsed = json.loads(txt)
        if isinstance(parsed, dict):
            return parsed
    except Exception:
        pass
    try:
        parsed = ast.literal_eval(txt)
        if isinstance(parsed, dict):
            return parsed
    except Exception:
        pass
    return {}


def _extract_rendered_doc_path(state: Optional[EngineState]) -> Optional[str]:
    if not state or not state.last_tool_events:
        return None
    for ev in reversed(state.last_tool_events):
        if ev.tool_name != "bua_render":
            continue
        payload = _parse_tool_result_payload(ev.response.get("result"))
        file_path = payload.get("file_path")
        if isinstance(file_path, str) and file_path.strip():
            return file_path.strip()
    return state.last_generated_doc_path


def _extract_latest_bua_state(state: Optional[EngineState]) -> Optional[dict[str, Any]]:
    if not state or not state.last_tool_events:
        return None
    for ev in reversed(state.last_tool_events):
        if ev.tool_name not in {"bua_questionnaire_data_parser", "questionnaire_parsing"}:
            continue
        payload = _parse_tool_result_payload(ev.response.get("result"))
        current_state = payload.get("current_state")
        if isinstance(current_state, dict):
            return current_state
    return None


def _extract_questionnaire_complete(state: Optional[EngineState]) -> bool:
    if not state or not state.last_tool_events:
        return False
    for ev in reversed(state.last_tool_events):
        if ev.tool_name not in {"bua_questionnaire_data_parser", "questionnaire_parsing"}:
            continue
        payload = _parse_tool_result_payload(ev.response.get("result"))
        if (
            payload.get("status") == "success"
            and payload.get("next_stage_to_fetch") == "complete"
        ):
            return True
    return False


def _build_fallback_assistant_text(state: Optional[EngineState]) -> str:
    if not state or not state.last_tool_events:
        return "_No text response._"
    for ev in reversed(state.last_tool_events):
        payload = _parse_tool_result_payload(ev.response.get("result"))
        if payload:
            prompt_text = payload.get("prompt_text")
            if isinstance(prompt_text, str) and prompt_text.strip():
                return prompt_text.strip()
            message = payload.get("message")
            if isinstance(message, str) and message.strip():
                return message.strip()
        err = ev.response.get("error")
        if isinstance(err, str) and err.strip():
            return f"Tool error from `{ev.tool_name}`: {err}"
    return "_No text response._"


def _render_bua_preview_panel() -> None:
    live_state = st.session_state.get("live_bua_state")
    with st.expander("BUA Preview", expanded=True):
        if not live_state:
            st.caption("No BUA state available yet.")
            return
        st.json(live_state, expanded=False)


def _run_turn(engine: MCPGeminiEngine, user_text: str) -> None:
    kickoff_text = "Start the BUA questionnaire."
    if user_text.strip() == kickoff_text:
        # Strong duplicate guard against Streamlit reruns/button replays.
        for m in st.session_state.get("messages", []):
            if m.get("role") == "user" and str(m.get("content", "")).strip() == kickoff_text:
                return
        if st.session_state.get("questionnaire_started"):
            return
        st.session_state.questionnaire_started = True

    st.session_state.messages.append({"role": "user", "content": user_text})
    with st.chat_message("user"):
        st.markdown(user_text)

    with st.chat_message("assistant"):
        with st.spinner("Thinking..."):
            try:
                reply, state = run_coroutine_sync(_send(engine, user_text))
            except asyncio.TimeoutError:
                reply = (
                    "Timed out waiting for a model/tool response. "
                    "The MCP tool call likely stalled and was aborted; please retry this step."
                )
                state = getattr(engine, "state", None)
            except Exception as e:
                reply = f"Error: {e}"
                state = getattr(engine, "state", None)

        rendered_path = _extract_rendered_doc_path(state)
        if rendered_path:
            st.session_state.last_doc_path = rendered_path
            p = Path(rendered_path)
            if p.exists():
                st.session_state.last_render_text = _docx_to_text(p)
        live_state = _extract_latest_bua_state(state)
        if live_state:
            st.session_state.live_bua_state = live_state
        if _extract_questionnaire_complete(state):
            st.session_state.questionnaire_complete = True

        assistant_text = (reply or "").strip()
        if not assistant_text:
            assistant_text = _build_fallback_assistant_text(state)
        st.markdown(assistant_text)
        st.session_state.messages.append({"role": "assistant", "content": assistant_text})


def _render_start_screen() -> None:
    st.subheader("START")
    c1, c2 = st.columns(2)
    with c1:
        st.markdown("### Start Fresh (questionaire)")
        st.caption("Begin interactive BUA question flow from scratch.")
        if st.button("Start Fresh (questionaire)", use_container_width=True):
            st.session_state.mode = "fresh"
            st.rerun()
    with c2:
        st.markdown("### upload file (docx)")
        st.caption("Upload an existing BUA draft and extract it into state.")
        if st.button("upload file (docx)", use_container_width=True):
            st.session_state.mode = "upload"
            st.rerun()


def _docx_to_text(doc_path: Path) -> str:
    try:
        doc = Document(str(doc_path))
        lines = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    cell_text = "\n".join(
                        p.text.strip() for p in cell.paragraphs if p.text and p.text.strip()
                    ).strip()
                    cells.append(cell_text)
                if any(cells):
                    lines.append(" | ".join(cells))
        text = "\n".join(lines).strip()
        if not text:
            return "(Rendered document has no plain text paragraphs to preview.)"
        if len(text) > 30000:
            return text[:30000] + "\n\n... [preview truncated]"
        return text
    except Exception as e:
        return f"(Could not preview rendered document text: {e})"


def _soffice_candidates() -> list[Path]:
    paths: list[Path] = []
    env_exe = os.environ.get("LIBREOFFICE_SOFFICE_PATH", "").strip()
    if env_exe:
        paths.append(Path(env_exe))
    paths.extend(
        [
            Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
            Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
        ]
    )
    out: list[Path] = []
    for raw in paths:
        p = Path(raw)
        if p.is_file():
            out.append(p)
    which = shutil.which("soffice")
    if which:
        wp = Path(which)
        out.append(wp)

    seen: set[str] = set()
    uniq: list[Path] = []
    for raw in out:
        try:
            key = str(raw.resolve())
        except Exception:
            key = str(raw)
        if key in seen:
            continue
        seen.add(key)
        uniq.append(raw)
    return uniq


def _docx_to_pdf_bytes(doc_path: Path) -> tuple[Optional[bytes], str]:
    """
    Produce PDF bytes preserving layout / table lines (unlike HTML conversion).
    Tries LibreOffice headless first, then docx2pdf (requires Microsoft Word on Windows/Mac).
    Optional env: DOCX_PREVIEW_PDF_CONVERTER=libreoffice|word|auto (default auto),
    LIBREOFFICE_SOFFICE_PATH=full path to soffice.exe
    """
    doc_path = doc_path.resolve()
    if not doc_path.is_file():
        return None, "Document file not found."

    mode = os.environ.get("DOCX_PREVIEW_PDF_CONVERTER", "auto").strip().lower()
    word_err = ""

    tmp_root = tempfile.mkdtemp(prefix="bua_pdf_")
    try:
        if mode in ("", "auto", "libreoffice"):
            for soffice in _soffice_candidates():
                cmd = [
                    str(soffice),
                    "--headless",
                    "--norestore",
                    "--nologo",
                    "--nofirststartwizard",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(tmp_root),
                    str(doc_path),
                ]
                try:
                    subprocess.run(
                        cmd,
                        check=False,
                        timeout=180,
                        capture_output=True,
                        text=True,
                    )
                except (FileNotFoundError, subprocess.TimeoutExpired, OSError):
                    continue

                pdf_file = Path(tmp_root) / f"{doc_path.stem}.pdf"
                if pdf_file.is_file() and pdf_file.stat().st_size > 400:
                    return pdf_file.read_bytes(), "LibreOffice"

        if mode in ("", "auto", "word"):
            out_pdf = Path(tmp_root) / f"{doc_path.stem}.pdf"
            try:
                from docx2pdf import convert as docx_convert  # type: ignore[import-untyped]

                docx_convert(str(doc_path), str(out_pdf))
                if out_pdf.is_file() and out_pdf.stat().st_size > 400:
                    return out_pdf.read_bytes(), "Microsoft Word (docx2pdf)"
            except Exception as e:
                word_err = str(e)

        suffix = f" Detail: {word_err[:280]}" if word_err else ""
        return (
            None,
            "Could not produce a PDF preview. Install LibreOffice (recommended) or "
            "`pip install docx2pdf` with Word installed." + suffix,
        )
    finally:
        shutil.rmtree(tmp_root, ignore_errors=True)


def _show_pdf_embedding(pdf_bytes: bytes, height: int = 720) -> None:
    show = getattr(st, "pdf", None)
    if callable(show):
        try:
            show(pdf_bytes, height=height)
            return
        except StreamlitAPIException:
            st.caption("Native PDF viewer unavailable (`pip install 'streamlit[pdf]'`). Using iframe fallback.")
        except Exception:
            st.caption("Could not render with st.pdf — using iframe fallback.")
    b64 = base64.b64encode(pdf_bytes).decode("utf-8")
    st_html(
        f'<iframe src="data:application/pdf;base64,{b64}" width="100%" height="{height}"></iframe>',
        height=height + 24,
        scrolling=True,
    )


def _docx_to_html(doc_path: Path) -> str:
    try:
        import mammoth

        with open(doc_path, "rb") as f:
            result = mammoth.convert_to_html(f)
        body = result.value or ""
        return (
            "<div style='font-family: Arial, sans-serif; padding: 0.75rem; line-height: 1.45;'>"
            f"{body}"
            "</div>"
        )
    except Exception as e:
        return (
            "<div style='font-family: Arial, sans-serif; padding: 0.75rem;'>"
            f"<p><strong>Visual preview unavailable.</strong> {e}</p>"
            "<p>Use the download button for the exact Word formatting.</p>"
            "</div>"
        )


def main() -> None:
    st.set_page_config(page_title="BUA Assistant", layout="wide")
    _init_state()
    engine = _get_engine()
    _render_sidebar_tool_library()

    st.title("Biological Use Authorization (BUA) Assistant")
    _render_debug_panel(getattr(engine, "state", None))

    if st.session_state.mode is None:
        _render_start_screen()
        return

    with st.sidebar:
        st.markdown("### Session")
        st.write(f"Mode: `{st.session_state.mode}`")
        if st.button("Back to START", use_container_width=True):
            st.session_state.mode = None
            st.session_state.upload_processed = False
            st.session_state.questionnaire_complete = False
            st.session_state.questionnaire_started = False
            st.session_state.live_bua_state = None
            st.session_state.messages = []
            st.rerun()

    if st.session_state.mode == "upload":
        uploads_dir = PROJECT_DIR / "uploads"
        uploads_dir.mkdir(parents=True, exist_ok=True)
        uploaded = st.file_uploader("Upload .docx", type=["docx"])
        if uploaded is not None:
            safe_name = Path(uploaded.name).name
            out_path = uploads_dir / safe_name
            out_path.write_bytes(uploaded.getbuffer())
            st.session_state.uploaded_path = str(out_path)
            st.success(f"Uploaded: {safe_name}")
            if not st.session_state.upload_processed:
                kickoff = (
                    "Parse this uploaded document and use it to populate BUA state. "
                    f"File path: {st.session_state.uploaded_path}. "
                    "Call doc_parsing, then save extracted data with bua_upload or stage parsers."
                )
                _run_turn(engine, kickoff)
                st.session_state.upload_processed = True

    if st.session_state.mode == "fresh" and not st.session_state.messages:
        if st.button("Start questionnaire now", type="primary"):
            st.session_state.questionnaire_complete = False
            st.session_state.questionnaire_started = False
            _run_turn(engine, "Start the BUA questionnaire.")

    for m in st.session_state.messages:
        with st.chat_message(m["role"]):
            st.markdown(m["content"])

    user_input = st.chat_input("Type your message and press Enter")
    if user_input:
        _run_turn(engine, user_input)

    if st.session_state.mode == "fresh" and st.session_state.questionnaire_complete:
        _render_bua_preview_panel()

    doc_path = st.session_state.get("last_doc_path")
    if doc_path:
        p = Path(doc_path)
        if p.exists():
            with st.expander("Rendered Document Viewer (PDF)", expanded=True):
                cache_key = ("pdf_preview", str(p.resolve()), p.stat().st_mtime_ns)
                if st.session_state.get("_render_pdf_cache_key") != cache_key:
                    pdf_bytes, pdf_note = _docx_to_pdf_bytes(p)
                    st.session_state["_render_pdf_cache_key"] = cache_key
                    st.session_state["_render_pdf_bytes"] = pdf_bytes
                    st.session_state["_render_pdf_note"] = pdf_note
                pdf_bytes = st.session_state.get("_render_pdf_bytes")
                pdf_note = str(st.session_state.get("_render_pdf_note") or "")

                if pdf_bytes:
                    st.caption(f"Layout-faithful preview ({pdf_note.strip()}).")
                    _show_pdf_embedding(pdf_bytes, height=720)
                    st.download_button(
                        "Download rendered PDF",
                        data=pdf_bytes,
                        file_name=p.with_suffix(".pdf").name,
                        mime="application/pdf",
                    )
                else:
                    st.warning(pdf_note or "PDF preview unavailable.")
                with st.expander("Approximate HTML preview (no boxes)", expanded=False):
                    st_html(_docx_to_html(p), height=480, scrolling=True)

            preview_text = st.session_state.get("last_render_text")
            if preview_text:
                with st.expander("Rendered Document Preview (plain text)", expanded=False):
                    st.text_area("Rendered BUA text", value=preview_text, height=320)
            with st.sidebar:
                st.success("Latest BUA document generated")
                with open(p, "rb") as f:
                    st.download_button(
                        "Download .docx",
                        data=f,
                        file_name=p.name,
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )

if __name__ == "__main__":
    main()
