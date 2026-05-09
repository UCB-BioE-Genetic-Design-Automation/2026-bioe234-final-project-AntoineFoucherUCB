"""Tests for `modules.biosafety.tools.doc_parsing` (uploaded .docx → markdown for the LLM)."""

from __future__ import annotations

from pathlib import Path

from modules.biosafety.tools import doc_parsing


def test_doc_parsing_missing_file(tmp_path: Path):
    missing = tmp_path / "does_not_exist.docx"
    out = doc_parsing.doc_parsing(str(missing))
    assert out["status"] == "error"
    assert str(missing) in out["message"]


def test_doc_parsing_success_minimal_docx(tmp_path: Path):
    docx = tmp_path / "tiny.docx"
    import docx as python_docx

    d = python_docx.Document()
    d.add_paragraph("BUA Title")
    d.add_paragraph("PI: Test User")
    d.save(docx)

    out = doc_parsing.doc_parsing(str(docx))
    assert out["status"] == "success"
    assert "BUA Title" in out["document_text"]
